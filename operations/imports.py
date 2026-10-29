import csv
import io
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from django.core.exceptions import ValidationError, PermissionDenied
from django.contrib.auth import get_user_model
from django.db import transaction
from openpyxl import load_workbook
from loans.models import Loan
from loans.utils import calculate_emi
from .models import BankEntry

LOAN_HEADERS = [
    "external_id",
    "loan_name",
    "loan_type",
    "amount",
    "interest_rate",
    "tenure_years",
    "start_date",
    "first_emi_date",
    "emi_frequency",
]
BANK_HEADERS = ["reference", "amount", "received_on"]


def read_rows(file):
    if file.size > 2 * 1024 * 1024:
        raise ValidationError("Maximum file size is 2 MB.")
    ext = file.name.lower().rsplit(".", 1)[-1]
    if ext == "csv":
        try:
            rows = list(csv.reader(io.StringIO(file.read().decode("utf-8-sig"))))
        except UnicodeDecodeError:
            raise ValidationError("CSV must use UTF-8 encoding.")
    elif ext == "xlsx":
        # Reject zip bombs before openpyxl reads workbook XML.
        import zipfile

        with zipfile.ZipFile(file) as archive:
            if sum(x.file_size for x in archive.infolist()) > 20 * 1024 * 1024:
                raise ValidationError("Expanded workbook exceeds 20 MB.")
        file.seek(0)
        wb = load_workbook(file, read_only=True, data_only=False)
        rows = []
        for row in wb.active.iter_rows(values_only=True):
            rows.append(list(row))
            if len(rows) > 1001:
                break
        wb.close()
    else:
        raise ValidationError("Choose a CSV or XLSX file.")
    if len(rows) < 2 or len(rows) > 1001:
        raise ValidationError("Provide between 1 and 1,000 data rows.")
    headers = [str(c or "").strip() for c in rows[0]]
    if len(headers) != len(set(headers)):
        raise ValidationError("Duplicate column headers.")
    return [
        dict(
            zip(
                headers,
                [
                    (
                        v.isoformat()[:10]
                        if isinstance(v, (date, datetime))
                        else str(v if v is not None else "").strip()
                    )
                    for v in row
                ],
            )
        )
        for row in rows[1:]
        if any(v is not None and str(v).strip() for v in row)
    ]


def decimal_value(value):
    result = Decimal(value)
    if (
        not result.is_finite()
        or result < 0
        or result != result.quantize(Decimal("0.01"))
    ):
        raise ValueError("Use non-negative amounts with at most two decimals.")
    return result


def validate_rows(rows, kind, user):
    errors = []
    seen = set()
    expected = LOAN_HEADERS if kind == "loans" else BANK_HEADERS
    for number, row in enumerate(rows, 2):
        try:
            if any(k not in row for k in expected):
                raise ValueError("Missing required columns: " + ", ".join(expected))
            key = row["external_id" if kind == "loans" else "reference"]
            if not key or len(key) > 100 or key in seen:
                raise ValueError("Missing, long or duplicate row identifier.")
            seen.add(key)
            for field in (
                ("external_id", "loan_name") if kind == "loans" else ("reference",)
            ):
                if row[field].lstrip().startswith(("=", "+", "-", "@")):
                    raise ValueError(
                        "Text identifiers must not start with spreadsheet formula characters."
                    )
            amount = decimal_value(row["amount"])
            if amount <= 0:
                raise ValueError("Amount must be positive.")
            if kind == "loans":
                import_key = f"{user.pk}:{key}"
                if Loan.objects.filter(import_key=import_key).exists():
                    raise ValueError("External ID already imported.")
                start = date.fromisoformat(row["start_date"])
                first = date.fromisoformat(row["first_emi_date"])
                if start > date.today() or first < start:
                    raise ValueError(
                        "Start must not be future; first EMI must not precede start."
                    )
                loan = Loan(
                    user=user,
                    import_key=import_key,
                    loan_name=row["loan_name"],
                    loan_type=row["loan_type"],
                    amount=amount,
                    interest_rate=decimal_value(row["interest_rate"]),
                    tenure_years=int(row["tenure_years"]),
                    start_date=start,
                    first_emi_date=first,
                    emi_frequency=row["emi_frequency"],
                    emi=0,
                    remaining_balance=0,
                )
                if not 1 <= loan.tenure_years <= 50 or loan.interest_rate > 50:
                    raise ValueError("Tenure must be 1–50 years and rate 0–50%.")
                loan.full_clean()
            else:
                if BankEntry.objects.filter(reference=key).exists():
                    raise ValueError("Bank reference already imported.")
                entry = BankEntry(
                    reference=key,
                    amount=amount,
                    received_on=date.fromisoformat(row["received_on"]),
                    uploaded_by=user,
                )
                if entry.received_on > date.today():
                    raise ValueError("Bank date cannot be future.")
                entry.full_clean()
        except (ValueError, InvalidOperation, ValidationError) as exc:
            errors.append(f"Row {number}: {exc}")
    return errors


@transaction.atomic
def commit_rows(rows, kind, user):
    errors = validate_rows(rows, kind, user)
    if errors:
        raise ValidationError(errors)
    for row in rows:
        if kind == "loans":
            loan = Loan(
                user=user,
                import_key=f"{user.pk}:{row['external_id']}",
                loan_name=row["loan_name"],
                loan_type=row["loan_type"],
                amount=decimal_value(row["amount"]),
                interest_rate=decimal_value(row["interest_rate"]),
                tenure_years=int(row["tenure_years"]),
                start_date=date.fromisoformat(row["start_date"]),
                first_emi_date=date.fromisoformat(row["first_emi_date"]),
                emi_frequency=row["emi_frequency"],
                remaining_balance=0,
                auto_debit=False,
            )
            loan.emi = calculate_emi(
                loan.amount, loan.interest_rate, loan.tenure_years, loan.emi_frequency
            )
            loan.save()
        else:
            BankEntry.objects.create(
                reference=row["reference"],
                amount=decimal_value(row["amount"]),
                received_on=date.fromisoformat(row["received_on"]),
                uploaded_by=user,
            )
    return len(rows)
